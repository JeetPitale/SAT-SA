"""
Official 6-Slide SIH 2026 Presentation Generator for SAT-SA
Strict adherence to the official 6-slide template format:
1. TITLE PAGE
2. IDEA TITLE
3. TECHNICAL APPROACH
4. FEASIBILITY AND VIABILITY
5. IMPACT AND BENEFITS
6. RESEARCH AND REFERENCES
"""

import os
import sys
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# Color Palette (SIH Official Theme)
C_NAVY_DARK = RGBColor(10, 25, 47)       # #0A192F (Deep Navy)
C_NAVY = RGBColor(15, 30, 60)            # #0F1E3C (Primary Navy)
C_NAVY_LIGHT = RGBColor(24, 43, 73)      # #182B49 (Dark Surface)
C_SIH_BLUE = RGBColor(0, 112, 192)       # #0070C0 (Official SIH Blue)
C_ACCENT_BLUE = RGBColor(2, 132, 199)    # #0284C7 (Cyan Blue)
C_SLATE_DARK = RGBColor(30, 41, 59)      # #1E293B (Headers)
C_SLATE_TEXT = RGBColor(51, 65, 85)      # #334155 (Body Text)
C_MUTED = RGBColor(100, 116, 139)        # #64748B (Muted)
C_BG_LIGHT = RGBColor(248, 250, 252)     # #F8FAFC (Light Canvas)
C_BG_CARD = RGBColor(241, 245, 249)      # #F1F5F9 (Card Surface)
C_BORDER = RGBColor(203, 213, 225)       # #CBD5E1 (Border)
C_BORDER_LIGHT = RGBColor(226, 232, 240) # #E2E8F0 (Soft Border)
C_WHITE = RGBColor(255, 255, 255)
C_AMBER = RGBColor(217, 119, 6)          # #D97706 (Warning Amber)
C_AMBER_BG = RGBColor(254, 243, 199)     # #FEF3C7 (Amber Container)
C_RED = RGBColor(220, 38, 38)            # #DC2626 (Critical Red)
C_RED_BG = RGBColor(254, 226, 226)       # #FEE2E2 (Red Container)
C_GREEN = RGBColor(5, 150, 105)          # #059669 (Healthy Green)
C_GREEN_BG = RGBColor(209, 250, 229)     # #D1FAE5 (Green Container)
C_BLUE_BG = RGBColor(224, 242, 254)      # #E0F2FE (Blue Container)

FONT_TITLE = "Arial"
FONT_BODY = "Arial"
FONT_SERIF = "Times New Roman"

def set_shape_style(shape, fill_color, border_color=None, border_width=1):
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(border_width)
    else:
        shape.line.fill.background()

def add_header(slide, slide_num, total_slides, sih_title, subhead=""):
    # Top-right SIH Logo
    logo_path = 'scratch/template_images/slide2_Picture 10_3.png'
    if os.path.exists(logo_path):
        slide.shapes.add_picture(logo_path, Inches(10.70), Inches(0.00), width=Inches(2.46), height=Inches(1.16))
    
    # Top-left Pill: Team Name / PS ID (matches template "Your Team Name" oval)
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.50), Inches(0.18), Inches(2.20), Inches(0.32))
    set_shape_style(pill, C_SIH_BLUE, None)
    tf = pill.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = "SIH 2026 | PS 26157"
    p.font.name = FONT_BODY
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.alignment = PP_ALIGN.CENTER
    
    # Official SIH Main Section Header
    title_box = slide.shapes.add_textbox(Inches(0.50), Inches(0.55), Inches(10.00), Inches(0.68))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
    
    p_title = tf_title.paragraphs[0]
    p_title.text = sih_title
    p_title.font.name = FONT_SERIF
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = C_NAVY

    if subhead:
        p_sub = tf_title.add_paragraph()
        p_sub.text = subhead
        p_sub.font.name = FONT_BODY
        p_sub.font.size = Pt(10.5)
        p_sub.font.bold = True
        p_sub.font.color.rgb = C_SIH_BLUE

    # Bottom blue ribbon
    ribbon = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.00), Inches(6.95), Inches(13.333), Inches(0.55))
    set_shape_style(ribbon, C_SIH_BLUE, None)
    
    # Footer text left
    foot_left = slide.shapes.add_textbox(Inches(0.50), Inches(7.02), Inches(5.20), Inches(0.40))
    tf_fl = foot_left.text_frame
    tf_fl.margin_left = tf_fl.margin_top = tf_fl.margin_right = tf_fl.margin_bottom = 0
    p_fl = tf_fl.paragraphs[0]
    p_fl.text = "SAT-SA: Security Audit & Supervisory Analytics"
    p_fl.font.name = FONT_BODY
    p_fl.font.size = Pt(11)
    p_fl.font.color.rgb = C_WHITE

    # Footer text center
    foot_mid = slide.shapes.add_textbox(Inches(5.08), Inches(7.02), Inches(4.50), Inches(0.40))
    tf_fm = foot_mid.text_frame
    tf_fm.margin_left = tf_fm.margin_top = tf_fm.margin_right = tf_fm.margin_bottom = 0
    p_fm = tf_fm.paragraphs[0]
    p_fm.text = "@SIH Idea submission- Template | NCIIPC / NTRO"
    p_fm.font.name = FONT_BODY
    p_fm.font.size = Pt(11)
    p_fm.font.color.rgb = C_WHITE
    p_fm.alignment = PP_ALIGN.CENTER

    # Slide number right
    foot_right = slide.shapes.add_textbox(Inches(11.50), Inches(7.02), Inches(1.33), Inches(0.40))
    tf_fr = foot_right.text_frame
    tf_fr.margin_left = tf_fr.margin_top = tf_fr.margin_right = tf_fr.margin_bottom = 0
    p_fr = tf_fr.paragraphs[0]
    p_fr.text = f"Slide {slide_num} of {total_slides}"
    p_fr.font.name = FONT_BODY
    p_fr.font.size = Pt(11)
    p_fr.font.bold = True
    p_fr.font.color.rgb = C_WHITE
    p_fr.alignment = PP_ALIGN.RIGHT

def add_card(slide, left, top, width, height, bg_color=C_BG_CARD, border_color=C_BORDER, border_width=1, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE):
    card = slide.shapes.add_shape(shape_type, Inches(left), Inches(top), Inches(width), Inches(height))
    set_shape_style(card, bg_color, border_color, border_width)
    return card

def add_callout(slide, left, top, width, height, text, icon="💡", bg_color=C_BLUE_BG, border_color=C_SIH_BLUE, text_color=C_NAVY, font_size=11, bold=True):
    box = add_card(slide, left, top, width, height, bg_color, border_color, border_width=1.5)
    tb = slide.shapes.add_textbox(Inches(left + 0.15), Inches(top + 0.08), Inches(width - 0.30), Inches(height - 0.16))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = f"{icon}  {text}"
    p.font.name = FONT_BODY
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = text_color
    p.alignment = PP_ALIGN.CENTER
    return box

def add_down_arrow(slide, left, top, width=0.22, height=0.22, color=C_SIH_BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(left), Inches(top), Inches(width), Inches(height))
    set_shape_style(shape, color, None)
    return shape

def generate_official_deck():
    prs = Presentation('SIH2026-IDEA-Presentation-Format.backup.pptx')
    blank_layout = prs.slide_layouts[6]
    
    TOTAL_SLIDES = 6
    
    # Trim to exactly 6 slides
    while len(prs.slides) > TOTAL_SLIDES:
        rId = prs.slides._sldIdLst[6].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[6]
        
    while len(prs.slides) < TOTAL_SLIDES:
        prs.slides.add_slide(blank_layout)

    # Clear all shapes
    for slide in prs.slides:
        for s in list(slide.shapes):
            sp = s._element
            sp.getparent().remove(sp)

    # =========================================================================
    # SLIDE 1: TITLE PAGE (Official Template Layout)
    # =========================================================================
    s1 = prs.slides[0]
    logo_path = 'scratch/template_images/slide1_Picture 1_2.png'
    if os.path.exists(logo_path):
        s1.shapes.add_picture(logo_path, Inches(10.70), Inches(0.00), width=Inches(2.46), height=Inches(1.16))
    
    tb_top = s1.shapes.add_textbox(Inches(0.50), Inches(0.18), Inches(9.50), Inches(0.40))
    p_tt = tb_top.text_frame.paragraphs[0]
    p_tt.text = "SMART INDIA HACKATHON 2026"
    p_tt.font.name = FONT_SERIF
    p_tt.font.size = Pt(20)
    p_tt.font.bold = True
    p_tt.font.color.rgb = C_SIH_BLUE

    tb_sub = s1.shapes.add_textbox(Inches(0.50), Inches(0.55), Inches(9.50), Inches(0.35))
    p_sub = tb_sub.text_frame.paragraphs[0]
    p_sub.text = "TITLE PAGE"
    p_sub.font.name = FONT_SERIF
    p_sub.font.size = Pt(13)
    p_sub.font.bold = True
    p_sub.font.color.rgb = C_NAVY

    # Hero Title Box
    tb_title = s1.shapes.add_textbox(Inches(0.50), Inches(0.92), Inches(10.00), Inches(1.35))
    tf_t = tb_title.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
    
    p_t1 = tf_t.paragraphs[0]
    p_t1.text = "SAT-SA: Security Audit & Supervisory Analytics"
    p_t1.font.name = FONT_TITLE
    p_t1.font.size = Pt(24)
    p_t1.font.bold = True
    p_t1.font.color.rgb = C_NAVY

    p_t2 = tf_t.add_paragraph()
    p_t2.text = "Turning SOC Records into Actionable Audit Evidence"
    p_t2.font.name = FONT_BODY
    p_t2.font.size = Pt(12.5)
    p_t2.font.bold = True
    p_t2.font.color.rgb = C_SIH_BLUE

    # Left Container: Metadata Form Fields
    add_card(s1, 0.50, 2.35, 6.20, 3.45, bg_color=C_BG_CARD, border_color=C_BORDER, border_width=1.5)
    tb_meta = s1.shapes.add_textbox(Inches(0.70), Inches(2.45), Inches(5.80), Inches(3.25))
    tf_m = tb_meta.text_frame
    tf_m.word_wrap = True
    
    meta_items = [
        ("Problem Statement ID –", " 26157 (SIH26157)"),
        ("Problem Statement Title –", " Supervisory Analytics Tool for SOC Assessment"),
        ("Theme –", " Blockchain & Cybersecurity"),
        ("PS Category –", " Software (100% Offline / Air-Gapped)"),
        ("Team ID –", " SIH26157_TEAM"),
        ("Team Name –", " SAT-SA Analytics (Registered on portal)"),
        ("Supervisory Target –", " NCIIPC / NTRO (Section 70A, IT Act 2000)")
    ]
    for i, (k, v) in enumerate(meta_items):
        p = tf_m.paragraphs[0] if i == 0 else tf_m.add_paragraph()
        p.space_after = Pt(4)
        r1 = p.add_run()
        r1.text = k
        r1.font.name = FONT_BODY
        r1.font.size = Pt(10)
        r1.font.bold = True
        r1.font.color.rgb = C_NAVY
        r2 = p.add_run()
        r2.text = f" {v}"
        r2.font.name = FONT_BODY
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = C_SLATE_TEXT

    # Right Container: 5-Step Concept Workflow Diagram
    add_card(s1, 6.90, 2.35, 5.93, 3.45, bg_color=C_NAVY_DARK, border_color=C_SIH_BLUE, border_width=1.5)
    tb_cw_t = s1.shapes.add_textbox(Inches(7.10), Inches(2.45), Inches(5.53), Inches(0.30))
    p_cwt = tb_cw_t.text_frame.paragraphs[0]
    p_cwt.text = "CORE SUPERVISORY WORKFLOW"
    p_cwt.font.name = FONT_TITLE
    p_cwt.font.size = Pt(11)
    p_cwt.font.bold = True
    p_cwt.font.color.rgb = C_SIH_BLUE

    concept_steps = [
        ("🗄️  1. SOC Records", "Alerts, Cases, Closures, Telemetry Dumps", C_NAVY_LIGHT),
        ("⚙️  2. SAT-SA Engine", "Offline Normalization & Data Quality Gate", C_SIH_BLUE),
        ("📊  3. Audit Analytics", "Rules, Peer Benchmarking & Anomaly ML", C_NAVY_LIGHT),
        ("📜  4. Forensic Evidence", "Explainable Dossiers & Peer Baseline Links", C_ACCENT_BLUE),
        ("👨‍💻  5. Human Review", "Targeted, Defensible Expert Investigation", C_GREEN)
    ]
    for idx, (ch, cd, cc) in enumerate(concept_steps):
        c_top = 2.85 + idx * 0.55
        add_card(s1, 7.10, c_top, 5.53, 0.44, bg_color=cc, border_color=C_ACCENT_BLUE if cc == C_NAVY_LIGHT else cc, border_width=1)
        tb_c = s1.shapes.add_textbox(Inches(7.20), Inches(c_top + 0.02), Inches(5.33), Inches(0.40))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True
        p = tf_c.paragraphs[0]
        r1 = p.add_run()
        r1.text = f"{ch} — "
        r1.font.name = FONT_BODY
        r1.font.size = Pt(9.0)
        r1.font.bold = True
        r1.font.color.rgb = C_WHITE
        r2 = p.add_run()
        r2.text = cd
        r2.font.name = FONT_BODY
        r2.font.size = Pt(8.0)
        r2.font.color.rgb = C_BORDER_LIGHT if cc in [C_NAVY_LIGHT, C_NAVY_DARK] else C_WHITE

    # Bottom Callout at Slide 1
    add_callout(s1, 0.50, 5.95, 12.33, 0.80, 
                "SOC Records  ➔  SAT-SA  ➔  Analytics  ➔  Evidence  ➔  Human Review\nHuman-in-the-Loop Audit Decision-Support System for National Critical Infrastructure Protection",
                icon="🛡️", bg_color=C_BLUE_BG, border_color=C_SIH_BLUE, text_color=C_NAVY, font_size=10.5, bold=True)

    ribbon1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.00), Inches(6.95), Inches(13.333), Inches(0.55))
    set_shape_style(ribbon1, C_SIH_BLUE, None)
    foot_s1 = s1.shapes.add_textbox(Inches(0.50), Inches(7.02), Inches(12.33), Inches(0.40))
    p_fs1 = foot_s1.text_frame.paragraphs[0]
    p_fs1.text = "SIH 2026 | Problem Statement 26157 | SAT-SA: Security Audit & Supervisory Analytics | Slide 1 of 6"
    p_fs1.font.name = FONT_BODY
    p_fs1.font.size = Pt(11)
    p_fs1.font.color.rgb = C_WHITE
    p_fs1.alignment = PP_ALIGN.CENTER


    # =========================================================================
    # SLIDE 2: IDEA TITLE / PROPOSED SOLUTION (3-Part Layout)
    # =========================================================================
    s2 = prs.slides[1]
    add_header(s2, 2, TOTAL_SLIDES, "IDEA TITLE", "SAT-SA: Security Audit & Supervisory Analytics")

    # Column 1: THE PROBLEM
    add_card(s2, 0.50, 1.35, 3.80, 4.45, bg_color=C_WHITE, border_color=C_RED, border_width=1.5)
    tb_p_h = s2.shapes.add_textbox(Inches(0.65), Inches(1.45), Inches(3.50), Inches(0.30))
    p_ph = tb_p_h.text_frame.paragraphs[0]
    p_ph.text = "1. THE PROBLEM"
    p_ph.font.name = FONT_TITLE
    p_ph.font.size = Pt(11)
    p_ph.font.bold = True
    p_ph.font.color.rgb = C_RED

    prob_cards = [
        ("High Data Volume", "SOCs generate millions of alerts, cases, closures, and asset records."),
        ("Limited Expert Review", "Auditors can only manually inspect < 0.1% random sample."),
        ("Missing Evidence", "Silent critical assets & dropped logs produce zero events to sample.")
    ]
    for idx, (ph, pd) in enumerate(prob_cards):
        p_top = 1.85 + idx * 1.00
        add_card(s2, 0.65, p_top, 3.50, 0.85, bg_color=C_RED_BG, border_color=C_RED, border_width=1)
        tb_p = s2.shapes.add_textbox(Inches(0.75), Inches(p_top + 0.05), Inches(3.30), Inches(0.75))
        tf_p = tb_p.text_frame
        tf_p.word_wrap = True
        p1 = tf_p.paragraphs[0]
        p1.text = ph
        p1.font.name = FONT_BODY
        p1.font.size = Pt(9.5)
        p1.font.bold = True
        p1.font.color.rgb = C_RED
        p2 = tf_p.add_paragraph()
        p2.text = pd
        p2.font.name = FONT_BODY
        p2.font.size = Pt(8.5)
        p2.font.color.rgb = C_SLATE_TEXT

    # Column 2: OUR SOLUTION
    add_card(s2, 4.50, 1.35, 4.33, 4.45, bg_color=C_NAVY_DARK, border_color=C_SIH_BLUE, border_width=1.5)
    tb_s_h = s2.shapes.add_textbox(Inches(4.65), Inches(1.45), Inches(4.03), Inches(0.30))
    p_sh = tb_s_h.text_frame.paragraphs[0]
    p_sh.text = "2. OUR PROPOSED SOLUTION"
    p_sh.font.name = FONT_TITLE
    p_sh.font.size = Pt(11)
    p_sh.font.bold = True
    p_sh.font.color.rgb = C_SIH_BLUE

    sol_nodes = [
        ("SOC Records", "Historical Alerts, Cases, Assets", C_NAVY_LIGHT),
        ("SAT-SA Analytics Engine", "• Rule Analysis  • Statistics (EB)\n• Peer Cohorts  • Anomaly ML", C_SIH_BLUE),
        ("Review Priorities", "Ranked Critical Sector Entities", C_AMBER),
        ("Human Expert Auditor", "Targeted Forensic Investigation", C_GREEN)
    ]
    for idx, (sh, sd, sc) in enumerate(sol_nodes):
        s_top = 1.85 + idx * 0.90
        s_hgt = 0.70 if idx == 1 else 0.55
        add_card(s2, 4.65, s_top, 4.03, s_hgt, bg_color=sc if sc != C_NAVY_LIGHT else C_NAVY_LIGHT, border_color=C_ACCENT_BLUE if sc == C_NAVY_LIGHT else sc, border_width=1)
        tb_s = s2.shapes.add_textbox(Inches(4.75), Inches(s_top + 0.04), Inches(3.83), Inches(s_hgt - 0.08))
        tf_s = tb_s.text_frame
        tf_s.word_wrap = True
        p1 = tf_s.paragraphs[0]
        p1.text = sh
        p1.font.name = FONT_BODY
        p1.font.size = Pt(9.5)
        p1.font.bold = True
        p1.font.color.rgb = C_WHITE if sc in [C_NAVY_LIGHT, C_SIH_BLUE, C_GREEN] else C_NAVY
        p2 = tf_s.add_paragraph()
        p2.text = sd
        p2.font.name = FONT_BODY
        p2.font.size = Pt(8.0)
        p2.font.color.rgb = C_BORDER_LIGHT if sc in [C_NAVY_LIGHT, C_SIH_BLUE, C_GREEN] else C_SLATE_DARK
        if idx < 3:
            add_down_arrow(s2, 6.55, s_top + s_hgt + 0.02, width=0.15, height=0.12, color=C_SIH_BLUE)

    # Column 3: INNOVATION
    add_card(s2, 9.03, 1.35, 3.80, 4.45, bg_color=C_WHITE, border_color=C_SIH_BLUE, border_width=1.5)
    tb_i_h = s2.shapes.add_textbox(Inches(9.18), Inches(1.45), Inches(3.50), Inches(0.30))
    p_ih = tb_i_h.text_frame.paragraphs[0]
    p_ih.text = "3. INNOVATION & UNIQUENESS"
    p_ih.font.name = FONT_TITLE
    p_ih.font.size = Pt(11)
    p_ih.font.bold = True
    p_ih.font.color.rgb = C_SIH_BLUE

    innov_cards = [
        ("01 — EXECUTION GAPS", "Evidence exists, but behaviour appears unusual (e.g. 4s fast closures, template notes).", C_RED_BG, C_RED),
        ("02 — NEGATIVE SPACE", "Expected evidence is missing or unusually sparse (e.g. silent SCADA assets, blackout shifts).", C_BLUE_BG, C_SIH_BLUE),
        ("03 — EVIDENCE TRAIL", "Every recommendation links to an unbroken chain of verifiable raw records.", C_GREEN_BG, C_GREEN)
    ]
    for idx, (ih, idesc, i_bg, i_col) in enumerate(innov_cards):
        i_top = 1.85 + idx * 1.00
        add_card(s2, 9.18, i_top, 3.50, 0.85, bg_color=i_bg, border_color=i_col, border_width=1)
        tb_i = s2.shapes.add_textbox(Inches(9.28), Inches(i_top + 0.05), Inches(3.30), Inches(0.75))
        tf_i = tb_i.text_frame
        tf_i.word_wrap = True
        p1 = tf_i.paragraphs[0]
        p1.text = ih
        p1.font.name = FONT_BODY
        p1.font.size = Pt(9.5)
        p1.font.bold = True
        p1.font.color.rgb = i_col
        p2 = tf_i.add_paragraph()
        p2.text = idesc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(8.5)
        p2.font.color.rgb = C_SLATE_TEXT

    # Bottom Callout Banner
    add_callout(s2, 0.50, 5.90, 12.33, 0.85,
                '"SAT-SA analyses not only what exists — but also what should exist."\nIdentifies where expert review should focus without replacing the human investigator.',
                icon="🎯", bg_color=C_BG_CARD, border_color=C_NAVY, text_color=C_NAVY, font_size=11, bold=True)


    # =========================================================================
    # SLIDE 3: TECHNICAL APPROACH (Workflow, Detection Engine & Tech Stack)
    # =========================================================================
    s3 = prs.slides[2]
    add_header(s3, 3, TOTAL_SLIDES, "TECHNICAL APPROACH", "Methodology, Workflow & Detection Engine")

    # Top: System Workflow Bar
    add_card(s3, 0.50, 1.35, 12.33, 0.70, bg_color=C_NAVY_DARK, border_color=C_SIH_BLUE, border_width=1.5)
    tb_wf_t = s3.shapes.add_textbox(Inches(0.65), Inches(1.38), Inches(12.03), Inches(0.25))
    p_wft = tb_wf_t.text_frame.paragraphs[0]
    p_wft.text = "COMPLETE SYSTEM WORKFLOW PIPELINE"
    p_wft.font.name = FONT_BODY
    p_wft.font.size = Pt(9.0)
    p_wft.font.bold = True
    p_wft.font.color.rgb = C_SIH_BLUE

    pipe_steps_s3 = ["DATA INPUT", "NORMALISATION", "FEATURE EXTRACTION", "DETECTION ENGINE", "PRIORITISATION", "EVIDENCE TRAIL", "HUMAN REVIEW"]
    for idx, ps in enumerate(pipe_steps_s3):
        p_left = 0.65 + idx * 1.73
        add_card(s3, p_left, 1.62, 1.62, 0.35, bg_color=C_NAVY_LIGHT, border_color=C_ACCENT_BLUE, border_width=1)
        tb_ps = s3.shapes.add_textbox(Inches(p_left), Inches(1.65), Inches(1.62), Inches(0.30))
        p = tb_ps.text_frame.paragraphs[0]
        p.text = ps
        p.font.name = FONT_BODY
        p.font.size = Pt(8.0)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.alignment = PP_ALIGN.CENTER

    # Middle Left: Data Input & Detection Engine
    add_card(s3, 0.50, 2.15, 6.00, 2.70, bg_color=C_WHITE, border_color=C_BORDER, border_width=1.5)
    tb_de_t = s3.shapes.add_textbox(Inches(0.65), Inches(2.22), Inches(5.70), Inches(0.25))
    p_det = tb_de_t.text_frame.paragraphs[0]
    p_det.text = "DATA INPUT & DETECTION ENGINE"
    p_det.font.name = FONT_TITLE
    p_det.font.size = Pt(10.5)
    p_det.font.bold = True
    p_det.font.color.rgb = C_NAVY

    tb_de_b = s3.shapes.add_textbox(Inches(0.65), Inches(2.50), Inches(5.70), Inches(2.25))
    tf_deb = tb_de_b.text_frame
    tf_deb.word_wrap = True
    
    de_points = [
        ("• Ingestion Sources:", " CSV | JSON | DB Exports (Alerts, Cases, Closures, Assets, Telemetry)"),
        ("• 4 Parallel Analytics Branches:", " Rules • Statistical Analysis (EB) • Peer Benchmarks • Anomaly ML"),
        ("• Dual-Paradigm Output:", " Splits into Execution Gaps (Faulty actions) + Negative Space (Omissions)")
    ]
    for i, (k, v) in enumerate(de_points):
        p = tf_deb.paragraphs[0] if i == 0 else tf_deb.add_paragraph()
        p.space_after = Pt(4)
        r1 = p.add_run()
        r1.text = k
        r1.font.name = FONT_BODY
        r1.font.size = Pt(9.0)
        r1.font.bold = True
        r1.font.color.rgb = C_SIH_BLUE
        r2 = p.add_run()
        r2.text = f" {v}"
        r2.font.name = FONT_BODY
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = C_SLATE_DARK

    # Middle Right: Visual Mini-Demos
    add_card(s3, 6.70, 2.15, 6.13, 2.70, bg_color=C_BG_CARD, border_color=C_BORDER, border_width=1.5)
    
    add_card(s3, 6.85, 2.25, 5.83, 1.15, bg_color=C_RED_BG, border_color=C_RED, border_width=1)
    tb_md1 = s3.shapes.add_textbox(Inches(6.95), Inches(2.28), Inches(5.63), Inches(1.05))
    tf_m1 = tb_md1.text_frame
    tf_m1.word_wrap = True
    p1 = tf_m1.paragraphs[0]
    p1.text = "EXECUTION GAP VISUAL (Alert A-19283)"
    p1.font.name = FONT_BODY
    p1.font.size = Pt(9.0)
    p1.font.bold = True
    p1.font.color.rgb = C_RED
    p2 = tf_m1.add_paragraph()
    p2.text = "Opened: 10:02:15  ➔  Closed: 10:02:19  |  Duration: 4 seconds"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(9.0)
    p2.font.bold = True
    p2.font.color.rgb = C_NAVY
    p3 = tf_m1.add_paragraph()
    p3.text = "Peer Deviation: -3.82σ  ➔  [ Review Recommended ]"
    p3.font.name = FONT_BODY
    p3.font.size = Pt(8.5)
    p3.font.bold = True
    p3.font.color.rgb = C_RED

    add_card(s3, 6.85, 3.55, 5.83, 1.15, bg_color=C_BLUE_BG, border_color=C_SIH_BLUE, border_width=1)
    tb_md2 = s3.shapes.add_textbox(Inches(6.95), Inches(3.58), Inches(5.63), Inches(1.05))
    tf_m2 = tb_md2.text_frame
    tf_m2.word_wrap = True
    p1 = tf_m2.paragraphs[0]
    p1.text = "NEGATIVE SPACE VISUAL (Critical SCADA Gateways)"
    p1.font.name = FONT_BODY
    p1.font.size = Pt(9.0)
    p1.font.bold = True
    p1.font.color.rgb = C_SIH_BLUE
    p2 = tf_m2.add_paragraph()
    p2.text = "Expected: ████████████████████ (100% Baseline)"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(8.0)
    p2.font.color.rgb = C_SLATE_TEXT
    p3 = tf_m2.add_paragraph()
    p3.text = "Observed: ██████ (28%)  ➔  [ Potential Coverage Gap ]"
    p3.font.name = FONT_BODY
    p3.font.size = Pt(8.5)
    p3.font.bold = True
    p3.font.color.rgb = C_AMBER

    # Bottom Tech Stack
    add_card(s3, 0.50, 4.95, 12.33, 0.85, bg_color=C_NAVY_DARK, border_color=C_SIH_BLUE, border_width=1.5)
    tb_tech = s3.shapes.add_textbox(Inches(0.65), Inches(5.00), Inches(12.03), Inches(0.75))
    tf_tech = tb_tech.text_frame
    tf_tech.word_wrap = True
    p_t1 = tf_tech.paragraphs[0]
    p_t1.text = "TECHNOLOGY STACK & OFFLINE-FIRST ARCHITECTURE"
    p_t1.font.name = FONT_BODY
    p_t1.font.size = Pt(9.0)
    p_t1.font.bold = True
    p_t1.font.color.rgb = C_SIH_BLUE
    
    p_t2 = tf_tech.add_paragraph()
    p_t2.text = "Python 3  •  Polars / Pandas  •  DuckDB + Parquet  •  scikit-learn  •  FastAPI  •  Web Dashboard (Streamlit/Plotly)"
    p_t2.font.name = FONT_BODY
    p_t2.font.size = Pt(8.5)
    p_t2.font.bold = True
    p_t2.font.color.rgb = C_WHITE

    p_t3 = tf_tech.add_paragraph()
    p_t3.text = "🔒 OFFLINE-FIRST: Air-Gapped • No External Cloud APIs • Bit-for-Bit Reproducible • Tamper-Evident SHA-256 Ledger"
    p_t3.font.name = FONT_BODY
    p_t3.font.size = Pt(8.0)
    p_t3.font.color.rgb = C_AMBER

    # Bottom Callout Banner
    add_callout(s3, 0.50, 5.90, 12.33, 0.85,
                "Processes 500,000+ records in <15s with DuckDB/Parquet in a 100% air-gapped environment with zero cloud dependency.",
                icon="⚡", bg_color=C_BLUE_BG, border_color=C_SIH_BLUE, text_color=C_NAVY, font_size=10.5, bold=True)


    # =========================================================================
    # SLIDE 4: FEASIBILITY AND VIABILITY
    # =========================================================================
    s4 = prs.slides[3]
    add_header(s4, 4, TOTAL_SLIDES, "FEASIBILITY AND VIABILITY", "Feasibility Analysis & Risk-Mitigation Matrix")

    # Left: Feasibility Cards
    add_card(s4, 0.50, 1.35, 5.00, 4.45, bg_color=C_WHITE, border_color=C_BORDER, border_width=1.5)
    tb_fs_t = s4.shapes.add_textbox(Inches(0.70), Inches(1.45), Inches(4.60), Inches(0.30))
    p_fst = tb_fs_t.text_frame.paragraphs[0]
    p_fst.text = "FEASIBILITY ANALYSIS"
    p_fst.font.name = FONT_TITLE
    p_fst.font.size = Pt(11)
    p_fst.font.bold = True
    p_fst.font.color.rgb = C_NAVY

    feas_cards = [
        ("🔒 OFFLINE DEPLOYMENT", "Works completely without cloud access, external APIs, or outbound network calls in air-gapped environments."),
        ("⚡ SCALABLE ANALYTICS", "Columnar storage (DuckDB/Parquet) and vectorized processing easily handle millions of records locally."),
        ("🧩 MODULAR DETECTORS", "Detector rules, statistics, and anomaly models can be independently updated and verified."),
        ("👨‍💻 HUMAN-IN-THE-LOOP", "System prioritises and recommends review; human experts retain absolute assessment authority.")
    ]
    for idx, (fh, fd) in enumerate(feas_cards):
        f_top = 1.85 + idx * 0.95
        add_card(s4, 0.70, f_top, 4.60, 0.82, bg_color=C_BG_CARD, border_color=C_BORDER_LIGHT, border_width=1)
        tb_f = s4.shapes.add_textbox(Inches(0.80), Inches(f_top + 0.05), Inches(4.40), Inches(0.72))
        tf_f = tb_f.text_frame
        tf_f.word_wrap = True
        p1 = tf_f.paragraphs[0]
        p1.text = fh
        p1.font.name = FONT_BODY
        p1.font.size = Pt(9.5)
        p1.font.bold = True
        p1.font.color.rgb = C_NAVY
        p2 = tf_f.add_paragraph()
        p2.text = fd
        p2.font.name = FONT_BODY
        p2.font.size = Pt(8.0)
        p2.font.color.rgb = C_SLATE_TEXT

    # Right: Challenges & Mitigation Table
    add_card(s4, 5.70, 1.35, 7.13, 3.25, bg_color=C_WHITE, border_color=C_BORDER, border_width=1.5)
    tb_cm_t = s4.shapes.add_textbox(Inches(5.90), Inches(1.45), Inches(6.70), Inches(0.28))
    p_cmt = tb_cm_t.text_frame.paragraphs[0]
    p_cmt.text = "CHALLENGES & SAT-SA MITIGATION STRATEGIES"
    p_cmt.font.name = FONT_TITLE
    p_cmt.font.size = Pt(11)
    p_cmt.font.bold = True
    p_cmt.font.color.rgb = C_NAVY

    tbl_shape = s4.shapes.add_table(6, 2, Inches(5.90), Inches(1.80), Inches(6.73), Inches(2.65))
    tbl = tbl_shape.table
    tbl.columns[0].width = Inches(2.60)
    tbl.columns[1].width = Inches(4.13)

    cm_data = [
        ["Potential Challenge", "SAT-SA Mitigation Strategy"],
        ["Incomplete / Missing data", "Automated data-quality gates & schema validation"],
        ["Different CSE data schemas", "Canonical Pydantic schema maps heterogeneous formats"],
        ["Small sample sizes per entity", "Empirical Bayes Poisson-Gamma statistical shrinkage"],
        ["False alarm fatigue", "Context-aware peer benchmarking + human review"],
        ["Adversarial KPI gaming", "Multi-detector diversity + random audit baseline"]
    ]
    for r_idx, row in enumerate(cm_data):
        for c_idx, val in enumerate(row):
            cell = tbl.cell(r_idx, c_idx)
            cell.text = val
            p = cell.text_frame.paragraphs[0]
            p.font.name = FONT_BODY
            p.font.size = Pt(8.5 if r_idx > 0 else 9.0)
            if r_idx == 0:
                p.font.bold = True
                p.font.color.rgb = C_WHITE
                cell.fill.solid()
                cell.fill.fore_color.rgb = C_NAVY
            else:
                p.font.bold = (c_idx == 0)
                p.font.color.rgb = C_NAVY if c_idx == 0 else C_SLATE_DARK
                cell.fill.solid()
                cell.fill.fore_color.rgb = C_WHITE if r_idx % 2 == 1 else C_BG_CARD

    # Right Bottom: Deployment Workflow
    add_card(s4, 5.70, 4.70, 7.13, 1.10, bg_color=C_NAVY_DARK, border_color=C_SIH_BLUE, border_width=1.5)
    tb_dep = s4.shapes.add_textbox(Inches(5.85), Inches(4.75), Inches(6.83), Inches(0.95))
    tf_dep = tb_dep.text_frame
    tf_dep.word_wrap = True
    p1 = tf_dep.paragraphs[0]
    p1.text = "OFFLINE DEPLOYMENT WORKFLOW"
    p1.font.name = FONT_BODY
    p1.font.size = Pt(9.0)
    p1.font.bold = True
    p1.font.color.rgb = C_SIH_BLUE
    p2 = tf_dep.add_paragraph()
    p2.text = "Offline Dataset  ➔  SAT-SA Engine  ➔  Local DuckDB  ➔  Local Dashboard  ➔  Expert Review"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(8.5)
    p2.font.bold = True
    p2.font.color.rgb = C_WHITE
    p3 = tf_dep.add_paragraph()
    p3.text = "✓ Zero dependency on external cloud services or Internet connectivity."
    p3.font.name = FONT_BODY
    p3.font.size = Pt(8.0)
    p3.font.color.rgb = C_AMBER

    # Bottom Callout Box
    add_callout(s4, 0.50, 5.90, 12.33, 0.85,
                "Feasible, scalable, and completely viable for immediate air-gapped deployment across NCIIPC supervisory workflows.",
                icon="✅", bg_color=C_GREEN_BG, border_color=C_GREEN, text_color=C_NAVY, font_size=11, bold=True)


    # =========================================================================
    # SLIDE 5: IMPACT AND BENEFITS
    # =========================================================================
    s5 = prs.slides[4]
    add_header(s5, 5, TOTAL_SLIDES, "IMPACT AND BENEFITS", "Supervisory Transformation & Strategic Value")

    # Left: Transformation Diagram
    add_card(s5, 0.50, 1.35, 5.40, 4.45, bg_color=C_WHITE, border_color=C_BORDER, border_width=1.5)
    tb_tr_t = s5.shapes.add_textbox(Inches(0.70), Inches(1.45), Inches(5.00), Inches(0.28))
    p_trt = tb_tr_t.text_frame.paragraphs[0]
    p_trt.text = "SUPERVISORY TRANSFORMATION"
    p_trt.font.name = FONT_TITLE
    p_trt.font.size = Pt(11)
    p_trt.font.bold = True
    p_trt.font.color.rgb = C_NAVY

    add_card(s5, 0.70, 1.80, 5.00, 1.05, bg_color=C_RED_BG, border_color=C_RED, border_width=1)
    tb_td = s5.shapes.add_textbox(Inches(0.80), Inches(1.85), Inches(4.80), Inches(0.95))
    tf_td = tb_td.text_frame
    tf_td.word_wrap = True
    p1 = tf_td.paragraphs[0]
    p1.text = "TODAY: MANUAL DATA OVERLOAD"
    p1.font.name = FONT_BODY
    p1.font.size = Pt(9.0)
    p1.font.bold = True
    p1.font.color.rgb = C_RED
    p2 = tf_td.add_paragraph()
    p2.text = "Large SOC Dataset  ➔  Manual Sampling (<0.1%)  ➔  Limited Review Capacity  ➔  Potentially Missed Patterns"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(8.0)
    p2.font.color.rgb = C_SLATE_TEXT

    add_card(s5, 1.70, 2.95, 3.00, 0.35, bg_color=C_SIH_BLUE, border_color=C_SIH_BLUE, border_width=0)
    tb_sc = s5.shapes.add_textbox(Inches(1.70), Inches(2.98), Inches(3.00), Inches(0.30))
    p_sc = tb_sc.text_frame.paragraphs[0]
    p_sc.text = "▼  SAT-SA ANALYTICS  ▼"
    p_sc.font.name = FONT_BODY
    p_sc.font.size = Pt(8.5)
    p_sc.font.bold = True
    p_sc.font.color.rgb = C_WHITE
    p_sc.alignment = PP_ALIGN.CENTER

    add_card(s5, 0.70, 3.40, 5.00, 1.15, bg_color=C_BLUE_BG, border_color=C_SIH_BLUE, border_width=1)
    tb_ft = s5.shapes.add_textbox(Inches(0.80), Inches(3.45), Inches(4.80), Inches(1.05))
    tf_ft = tb_ft.text_frame
    tf_ft.word_wrap = True
    p1 = tf_ft.paragraphs[0]
    p1.text = "FUTURE: EVIDENCE-DRIVEN REVIEW"
    p1.font.name = FONT_BODY
    p1.font.size = Pt(9.0)
    p1.font.bold = True
    p1.font.color.rgb = C_SIH_BLUE
    p2 = tf_ft.add_paragraph()
    p2.text = "Evidence-Based Priorities  ➔  Explainable Findings  ➔  Focused Expert Review  ➔  Better Supervisory Assessment"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(8.0)
    p2.font.color.rgb = C_NAVY

    # Right Top: 4 Impact Cards
    add_card(s5, 6.10, 1.35, 6.73, 2.30, bg_color=C_WHITE, border_color=C_BORDER, border_width=1.5)
    tb_ic_t = s5.shapes.add_textbox(Inches(6.30), Inches(1.42), Inches(6.33), Inches(0.25))
    p_ict = tb_ic_t.text_frame.paragraphs[0]
    p_ict.text = "FOUR PILLARS OF SUPERVISORY VALUE"
    p_ict.font.name = FONT_TITLE
    p_ict.font.size = Pt(10.5)
    p_ict.font.bold = True
    p_ict.font.color.rgb = C_NAVY

    impact_pillars = [
        ("🔍 FOCUSED REVIEW", "Helps experts identify where to investigate first."),
        ("📊 EVIDENCE-DRIVEN", "Links recommendations directly to underlying records."),
        ("🛡️ BETTER VISIBILITY", "Detects both unusual behaviour and missing evidence."),
        ("⚙️ EFFICIENT AUDIT", "Reduces dependence on purely manual sampling.")
    ]
    for idx, (ih, idesc) in enumerate(impact_pillars):
        ix = 6.30 + (idx % 2) * 3.20
        iy = 1.72 + (idx // 2) * 0.90
        add_card(s5, ix, iy, 3.05, 0.82, bg_color=C_BG_CARD, border_color=C_BORDER_LIGHT, border_width=1)
        tb_p = s5.shapes.add_textbox(Inches(ix + 0.08), Inches(iy + 0.04), Inches(2.89), Inches(0.74))
        tf_p = tb_p.text_frame
        tf_p.word_wrap = True
        p1 = tf_p.paragraphs[0]
        p1.text = ih
        p1.font.name = FONT_BODY
        p1.font.size = Pt(9.0)
        p1.font.bold = True
        p1.font.color.rgb = C_NAVY
        p2 = tf_p.add_paragraph()
        p2.text = idesc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(8.0)
        p2.font.color.rgb = C_SLATE_TEXT

    # Right Bottom: Comparison Table
    add_card(s5, 6.10, 3.75, 6.73, 2.05, bg_color=C_NAVY_DARK, border_color=C_SIH_BLUE, border_width=1.5)
    tb_uv_t = s5.shapes.add_textbox(Inches(6.30), Inches(3.82), Inches(6.33), Inches(0.25))
    p_uvt = tb_uv_t.text_frame.paragraphs[0]
    p_uvt.text = "TRADITIONAL MANUAL REVIEW VS SAT-SA"
    p_uvt.font.name = FONT_TITLE
    p_uvt.font.size = Pt(10.5)
    p_uvt.font.bold = True
    p_uvt.font.color.rgb = C_SIH_BLUE

    uv_table_shape = s5.shapes.add_table(5, 2, Inches(6.30), Inches(4.12), Inches(6.33), Inches(1.55))
    tbl_uv = uv_table_shape.table
    tbl_uv.columns[0].width = Inches(3.15)
    tbl_uv.columns[1].width = Inches(3.18)

    uv_data = [
        ["Traditional Manual Review", "SAT-SA Supervisory Analytics"],
        ["Sample-driven (<0.1% random)", "Priority-driven (100% data coverage)"],
        ["Record-focused in isolation", "Behaviour-focused across time"],
        ["Hard to compare peer entities", "Context-aware peer benchmarking"],
        ["Inspects existing logs only", "Inspects existing + missing evidence"]
    ]
    for r_idx, row in enumerate(uv_data):
        for c_idx, val in enumerate(row):
            cell = tbl_uv.cell(r_idx, c_idx)
            cell.text = val
            p = cell.text_frame.paragraphs[0]
            p.font.name = FONT_BODY
            p.font.size = Pt(8.0 if r_idx > 0 else 8.5)
            if r_idx == 0:
                p.font.bold = True
                p.font.color.rgb = C_WHITE
                cell.fill.solid()
                cell.fill.fore_color.rgb = C_SIH_BLUE if c_idx == 1 else C_NAVY_LIGHT
            else:
                p.font.bold = (c_idx == 1)
                p.font.color.rgb = C_AMBER if c_idx == 1 else C_BORDER_LIGHT
                cell.fill.solid()
                cell.fill.fore_color.rgb = C_NAVY_LIGHT if r_idx % 2 == 1 else C_NAVY_DARK

    # Bottom Callout Box
    add_callout(s5, 0.50, 5.90, 12.33, 0.85,
                '"From Data Overload → Evidence-Driven Review"\nSAT-SA helps supervisory experts spend limited review time where evidence indicates the greatest need.',
                icon="🚀", bg_color=C_BLUE_BG, border_color=C_SIH_BLUE, text_color=C_NAVY, font_size=11, bold=True)


    # =========================================================================
    # SLIDE 6: RESEARCH AND REFERENCES
    # =========================================================================
    s6 = prs.slides[5]
    add_header(s6, 6, TOTAL_SLIDES, "RESEARCH  AND REFERENCES", "Regulatory Guidelines, Standards & Scientific References")

    ref_cards = [
        ("1. NCIIPC / GOVERNMENT OF INDIA GUIDELINES", [
            "• National Critical Information Infrastructure Protection Centre (NCIIPC) Guidelines (Section 70A, IT Act 2000).",
            "• NCIIPC / QCI Conformity Assessment Framework for Cybersecurity Auditing of Critical Sector Entities (CSEs).",
            "• Indian Computer Emergency Response Team (CERT-In) Cyber Security Directions for Incident Reporting & Triage."
        ], C_SIH_BLUE, C_BLUE_BG),
        ("2. SMART INDIA HACKATHON 2026 SPECIFICATION", [
            "• Problem Statement ID: 26157 (SIH26157) — Supervisory Analytics Tool for SOC Assessment.",
            "• Objective: Offline supervisory audit engine for prioritizing human investigation across Critical Sector Entity SOCs.",
            "• Operational Scope: Meta-analysis of historical alerts, case tickets, shift rosters, and asset telemetry inventories."
        ], C_NAVY, C_BG_CARD),
        ("3. RESEARCH ON SOC ANALYTICS & OPERATIONS", [
            "• MITRE ATT&CK Framework: Design and implementation of structured threat detection taxonomies for enterprise SOCs.",
            "• SOC Alert Fatigue and Supervisory Triage Optimization: IEEE/ACM Cybersecurity Surveys on Operational Bottlenecks.",
            "• Empirical Analysis of Security Operations Centers: Investigating closure velocity, shift drift, and ticket resolution quality."
        ], C_SLATE_DARK, C_BG_CARD),
        ("4. STATISTICAL & MACHINE LEARNING REFERENCES", [
            "• Isolation Forest for Unsupervised Anomaly Detection: Liu, Ting, & Zhou (ACM Transactions on Knowledge Discovery).",
            "• MinHash & Locality-Sensitive Hashing (LSH) for Near-Duplicate Text Fingerprinting: Broder et al.",
            "• Empirical Bayes Poisson-Gamma Shrinkage for Sparse Rate Stabilization: Efron & Morris (Journal of the American Statistical Association).",
            "• Information Retrieval Evaluation Metrics: Precision@K, Recall@K, and Detection Lift."
        ], C_GREEN, C_GREEN_BG)
    ]

    for idx, (rh, r_bullets, r_col, r_bg) in enumerate(ref_cards):
        rx = 0.50 + (idx % 2) * 6.38
        ry = 1.35 + (idx // 2) * 2.20
        add_card(s6, rx, ry, 5.95, 2.05, bg_color=C_WHITE, border_color=r_col, border_width=1.5)
        add_card(s6, rx, ry, 5.95, 0.40, bg_color=r_col, border_color=r_col, border_width=0)
        tb_rh = s6.shapes.add_textbox(Inches(rx + 0.15), Inches(ry + 0.05), Inches(5.65), Inches(0.30))
        p_rh = tb_rh.text_frame.paragraphs[0]
        p_rh.text = rh
        p_rh.font.name = FONT_BODY
        p_rh.font.size = Pt(9.5)
        p_rh.font.bold = True
        p_rh.font.color.rgb = C_WHITE

        tb_rb = s6.shapes.add_textbox(Inches(rx + 0.15), Inches(ry + 0.45), Inches(5.65), Inches(1.50))
        tf_rb = tb_rb.text_frame
        tf_rb.word_wrap = True
        for i, b in enumerate(r_bullets):
            p = tf_rb.paragraphs[0] if i == 0 else tf_rb.add_paragraph()
            p.space_after = Pt(2)
            p.text = b
            p.font.name = FONT_BODY
            p.font.size = Pt(8.0)
            p.font.color.rgb = C_SLATE_DARK

    # Bottom Callout Banner
    add_callout(s6, 0.50, 5.90, 12.33, 0.85,
                "All methodologies grounded in established supervisory standards, peer-reviewed statistical models, and official NCIIPC guidelines.",
                icon="📚", bg_color=C_BLUE_BG, border_color=C_SIH_BLUE, text_color=C_NAVY, font_size=11, bold=True)

    # Save presentation
    output_path = 'SIH2026-IDEA-Presentation-Format.pptx'
    prs.save(output_path)
    print(f"Official 6-Slide presentation successfully generated and saved to: {output_path}")

if __name__ == "__main__":
    generate_official_deck()
