"""
Comprehensive SAT-SA Presentation Builder for SIH 2026
Builds 11 slides using python-pptx with professional typography, cards, and diagrams.
"""

import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# Color Palette
C_NAVY_DARK = RGBColor(10, 25, 47)       # #0A192F
C_NAVY = RGBColor(15, 30, 60)            # #0F1E3C
C_NAVY_LIGHT = RGBColor(24, 43, 73)      # #182B49
C_SIH_BLUE = RGBColor(0, 112, 192)       # #0070C0 (Official template blue)
C_ACCENT_BLUE = RGBColor(2, 132, 199)    # #0284C7
C_SLATE_DARK = RGBColor(30, 41, 59)      # #1E293B
C_SLATE_TEXT = RGBColor(51, 65, 85)      # #334155
C_MUTED = RGBColor(100, 116, 139)        # #64748B
C_BG_LIGHT = RGBColor(248, 250, 252)     # #F8FAFC
C_BG_CARD = RGBColor(241, 245, 249)      # #F1F5F9
C_BORDER = RGBColor(203, 213, 225)       # #CBD5E1
C_BORDER_LIGHT = RGBColor(226, 232, 240) # #E2E8F0
C_WHITE = RGBColor(255, 255, 255)
C_AMBER = RGBColor(217, 119, 6)          # #D97706
C_AMBER_BG = RGBColor(254, 243, 199)     # #FEF3C7
C_RED = RGBColor(220, 38, 38)            # #DC2626
C_RED_BG = RGBColor(254, 226, 226)       # #FEE2E2
C_GREEN = RGBColor(5, 150, 105)          # #059669
C_GREEN_BG = RGBColor(209, 250, 229)     # #D1FAE5
C_BLUE_BG = RGBColor(224, 242, 254)      # #E0F2FE

FONT_TITLE = "Arial"
FONT_BODY = "Arial"

def set_flat_style(shape, fill_color, border_color=None, border_width=1):
    """Sets solid fill and border on a shape."""
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(border_width)
    else:
        shape.line.fill.background()

def add_header(slide, slide_num, total_slides, title_text, category_text="SECURITY AUDIT & SUPERVISORY ANALYTICS"):
    """Adds top-right SIH logo, bottom blue ribbon, footer metadata, team pill, and slide title."""
    # Top-right SIH Logo
    logo_path = 'scratch/template_images/slide2_Picture 10_3.png'
    if os.path.exists(logo_path):
        slide.shapes.add_picture(logo_path, Inches(10.70), Inches(0.00), width=Inches(2.46), height=Inches(1.16))
    
    # Top-left Team / PS pill
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.50), Inches(0.18), Inches(2.20), Inches(0.32))
    set_flat_style(pill, C_SIH_BLUE, None)
    tf = pill.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = "SIH 2026 | PS SIH26157"
    p.font.name = FONT_BODY
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.alignment = PP_ALIGN.CENTER
    
    # Category tag next to pill
    cat_box = slide.shapes.add_textbox(Inches(2.80), Inches(0.18), Inches(7.50), Inches(0.32))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = f"SAT-SA • {category_text.upper()}"
    p_cat.font.name = FONT_BODY
    p_cat.font.size = Pt(10)
    p_cat.font.bold = True
    p_cat.font.color.rgb = C_MUTED

    # Title text
    if title_text:
        title_box = slide.shapes.add_textbox(Inches(0.50), Inches(0.58), Inches(10.00), Inches(0.65))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = FONT_TITLE
        p_title.font.size = Pt(21)
        p_title.font.bold = True
        p_title.font.color.rgb = C_NAVY

    # Bottom blue ribbon
    ribbon = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.00), Inches(6.95), Inches(13.333), Inches(0.55))
    set_flat_style(ribbon, C_SIH_BLUE, None)
    
    # Footer text left
    foot_left = slide.shapes.add_textbox(Inches(0.50), Inches(7.02), Inches(5.00), Inches(0.40))
    tf_fl = foot_left.text_frame
    tf_fl.margin_left = tf_fl.margin_top = tf_fl.margin_right = tf_fl.margin_bottom = 0
    p_fl = tf_fl.paragraphs[0]
    p_fl.text = "SAT-SA: Supervisory Analytics Tool for SOC Assessment"
    p_fl.font.name = FONT_BODY
    p_fl.font.size = Pt(11)
    p_fl.font.color.rgb = C_WHITE

    # Footer text center
    foot_mid = slide.shapes.add_textbox(Inches(5.08), Inches(7.02), Inches(4.50), Inches(0.40))
    tf_fm = foot_mid.text_frame
    tf_fm.margin_left = tf_fm.margin_top = tf_fm.margin_right = tf_fm.margin_bottom = 0
    p_fm = tf_fm.paragraphs[0]
    p_fm.text = "@SIH Idea submission- Template | NCIIPC / NTRO"
    p_fm.font.name = FONT_BODY
    p_fm.font.size = Pt(11)
    p_fm.font.color.rgb = C_WHITE
    p_fm.alignment = PP_ALIGN.CENTER

    # Slide number right
    foot_right = slide.shapes.add_textbox(Inches(11.50), Inches(7.02), Inches(1.33), Inches(0.40))
    tf_fr = foot_right.text_frame
    tf_fr.margin_left = tf_fr.margin_top = tf_fr.margin_right = tf_fr.margin_bottom = 0
    p_fr = tf_fr.paragraphs[0]
    p_fr.text = f"Slide {slide_num} of {total_slides}"
    p_fr.font.name = FONT_BODY
    p_fr.font.size = Pt(11)
    p_fr.font.bold = True
    p_fr.font.color.rgb = C_WHITE
    p_fr.alignment = PP_ALIGN.RIGHT

def add_card(slide, left, top, width, height, bg_color=C_BG_CARD, border_color=C_BORDER, border_width=1, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE):
    """Creates a background container card."""
    card = slide.shapes.add_shape(shape_type, Inches(left), Inches(top), Inches(width), Inches(height))
    set_flat_style(card, bg_color, border_color, border_width)
    return card

def add_kpi_card(slide, left, top, width, height, value_text, label_text, sublabel="", val_color=C_SIH_BLUE, bg_color=C_WHITE, border_color=C_BORDER):
    """Adds an enterprise KPI card with metric and label."""
    card = add_card(slide, left, top, width, height, bg_color, border_color)
    tb = slide.shapes.add_textbox(Inches(left + 0.08), Inches(top + 0.08), Inches(width - 0.16), Inches(height - 0.16))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p1 = tf.paragraphs[0]
    p1.text = value_text
    p1.font.name = FONT_TITLE
    p1.font.size = Pt(22)
    p1.font.bold = True
    p1.font.color.rgb = val_color
    p1.alignment = PP_ALIGN.CENTER
    
    p2 = tf.add_paragraph()
    p2.text = label_text
    p2.font.name = FONT_BODY
    p2.font.size = Pt(10)
    p2.font.bold = True
    p2.font.color.rgb = C_SLATE_DARK
    p2.alignment = PP_ALIGN.CENTER
    
    if sublabel:
        p3 = tf.add_paragraph()
        p3.text = sublabel
        p3.font.name = FONT_BODY
        p3.font.size = Pt(8.5)
        p3.font.color.rgb = C_MUTED
        p3.alignment = PP_ALIGN.CENTER
    return card

def add_callout(slide, left, top, width, height, text, icon="💡", bg_color=C_BLUE_BG, border_color=C_SIH_BLUE, text_color=C_NAVY, bold=True):
    """Adds a callout quote / highlight box."""
    box = add_card(slide, left, top, width, height, bg_color, border_color, border_width=1.5)
    tb = slide.shapes.add_textbox(Inches(left + 0.15), Inches(top + 0.08), Inches(width - 0.30), Inches(height - 0.16))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = f"{icon}  {text}"
    p.font.name = FONT_BODY
    p.font.size = Pt(11)
    p.font.bold = bold
    p.font.color.rgb = text_color
    p.alignment = PP_ALIGN.CENTER
    return box

print("Core UI functions ready.")
